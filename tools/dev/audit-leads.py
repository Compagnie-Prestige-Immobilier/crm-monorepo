#!/usr/bin/env python3
"""Qualite du classeur de leads : ce qui est livre, ce qui est appelable.

Lit un classeur SharePoint (un onglet par jour) ou des CSV et compte, par jour
puis en cumul, les doublons et les lignes inexploitables. Le chiffre annonce
par le marketing est le nombre de lignes ; le chiffre utile est le nombre de
numeros uniques joignables. Un numero etranger est appelable : il compte comme
les autres, la colonne « etranger » ne fait que le signaler.

    python3 tools/dev/audit-leads.py "Leads du 10 sept 2026.xlsx"
    python3 tools/dev/audit-leads.py *.csv --csv rapport.csv
"""

import csv
import re
import sys
import unicodedata
from collections import Counter, OrderedDict

SCIENTIFIQUE = re.compile(r"^[+-]?\d(?:\.\d+)?[Ee][+-]?\d+$")
SEPARATEURS = re.compile(r"[\s.\-()/+]|^p:", re.IGNORECASE)
INDICATIFS_MOBILES_SN = ("70", "75", "76", "77", "78")
EMAIL = re.compile(r"^[^@\s]+@[^@\s]+\.[a-z]{2,}$", re.IGNORECASE)


def cle_entete(valeur):
    sans_accent = unicodedata.normalize("NFKD", str(valeur or ""))
    sans_accent = "".join(c for c in sans_accent if not unicodedata.combining(c))
    return re.sub(r"[^a-z]", "", sans_accent.lower())


def entier_lisible(valeur):
    """« 2.21773E+11 » : Excel a range le numero comme un nombre, il est tronque."""
    if not SCIENTIFIQUE.match(valeur):
        return valeur, False
    try:
        nombre = float(valeur)
    except ValueError:
        return valeur, False
    if nombre != int(nombre) or abs(nombre) >= 1e18:
        return valeur, False
    return str(int(nombre)), True


def normaliser_telephone(brut):
    """Rend (numero, etat) : etat vaut ok, tronque, etranger, illisible, vide."""
    brut = (brut or "").strip()
    if not brut:
        return "", "vide"
    valeur, scientifique = entier_lisible(brut)
    compact = SEPARATEURS.sub("", valeur)
    if compact.startswith("00"):
        compact = compact[2:]
    if not compact.isdigit():
        return "", "illisible"
    if scientifique:
        return compact, "tronque"
    if len(compact) == 9 and compact[:2] in INDICATIFS_MOBILES_SN:
        return "221" + compact, "ok"
    if compact.startswith("221"):
        return (compact, "ok") if len(compact) == 12 else (compact, "illisible")
    if 10 <= len(compact) <= 15:
        return compact, "etranger"
    return compact, "illisible"


def lire_csv(chemin):
    for encodage in ("utf-8-sig", "cp1252"):
        try:
            with open(chemin, newline="", encoding=encodage) as f:
                texte = f.read()
            break
        except UnicodeDecodeError:
            continue
    separateur = ";" if texte.count(";") > texte.count(",") else ","
    lignes = list(csv.reader(texte.splitlines(), delimiter=separateur))
    nom = chemin.rsplit("/", 1)[-1].rsplit(".", 1)[0]
    return [(nom, lignes)]


def lire_xlsx(chemin):
    from openpyxl import load_workbook

    classeur = load_workbook(chemin, read_only=True, data_only=True)
    feuilles = []
    for feuille in classeur.worksheets:
        lignes = [["" if c is None else str(c) for c in ligne] for ligne in feuille.iter_rows(values_only=True)]
        feuilles.append((feuille.title, lignes))
    return feuilles


def colonnes(entetes):
    """« Canal » decide du projet, « Provenance » n'est qu'un repli : voir
    docs/decisions/import-leads.md."""
    cibles = {"telephone": ("telephone", "tel", "phone", "numero"),
              "nom": ("nomcomplet", "nom", "name"),
              "email": ("email", "mail", "courriel"),
              "canal": ("canal", "provenance", "source")}
    trouvees = {}
    cles = [cle_entete(e) for e in entetes]
    for champ, noms in cibles.items():
        for nom in noms:
            if nom in cles:
                trouvees[champ] = cles.index(nom)
                break
    return trouvees


def auditer(feuilles):
    vus_telephone, vus_email = {}, {}
    rapport = OrderedDict()
    canaux = Counter()
    for nom, lignes in feuilles:
        entete = next((i for i, l in enumerate(lignes) if colonnes(l).get("telephone") is not None), None)
        if entete is None:
            continue
        cols = colonnes(lignes[entete])
        jour = {k: 0 for k in ("lignes", "ok", "vide", "illisible", "tronque", "etranger",
                               "doublon_jour", "doublon_anterieur", "doublon_email", "sans_nom")}
        dans_le_jour = set()
        for numero_ligne, ligne in enumerate(lignes[entete + 1:], start=entete + 2):
            def cellule(champ):
                i = cols.get(champ)
                return (ligne[i].strip() if i is not None and i < len(ligne) else "")
            if not any(cellule(c) for c in ("telephone", "nom", "email")):
                continue
            jour["lignes"] += 1
            if not cellule("nom"):
                jour["sans_nom"] += 1
            if cellule("canal"):
                canaux[cellule("canal")] += 1
            email = cellule("email").lower()
            if EMAIL.match(email):
                if email in vus_email:
                    jour["doublon_email"] += 1
                else:
                    vus_email[email] = (nom, numero_ligne)
            telephone, etat = normaliser_telephone(cellule("telephone"))
            if etat == "etranger":
                jour["etranger"] += 1
            elif etat != "ok":
                jour[etat] += 1
                continue
            if telephone in dans_le_jour:
                jour["doublon_jour"] += 1
            elif telephone in vus_telephone:
                jour["doublon_anterieur"] += 1
                dans_le_jour.add(telephone)
            else:
                jour["ok"] += 1
                vus_telephone[telephone] = (nom, numero_ligne)
                dans_le_jour.add(telephone)
        rapport[nom] = jour
    return rapport, canaux, len(vus_telephone)


COLONNES_RAPPORT = [("lignes", "livrees"), ("ok", "nettes"), ("doublon_jour", "dbl jour"),
                    ("doublon_anterieur", "dbl avant"), ("vide", "sans tel"),
                    ("illisible", "tel faux"), ("tronque", "tronque"),
                    ("etranger", "etranger"), ("doublon_email", "dbl mail"), ("sans_nom", "sans nom")]


def afficher(rapport, canaux, uniques):
    largeur = max([len(n) for n in rapport] + [12])
    entete = f"{'onglet'.ljust(largeur)}" + "".join(f"{t:>11}" for _, t in COLONNES_RAPPORT)
    print(entete)
    print("-" * len(entete))
    totaux = Counter()
    for nom, jour in rapport.items():
        print(nom.ljust(largeur) + "".join(f"{jour[c]:>11}" for c, _ in COLONNES_RAPPORT))
        totaux.update(jour)
    print("-" * len(entete))
    print("TOTAL".ljust(largeur) + "".join(f"{totaux[c]:>11}" for c, _ in COLONNES_RAPPORT))
    livrees, nettes = totaux["lignes"], totaux["ok"]
    print()
    print(f"UTILISABLE   {nettes:>6}   numeros uniques appelables"
          + (f"   {100 * nettes / livrees:.1f} %" if livrees else ""))
    causes = [("deja livre un jour precedent", totaux["doublon_anterieur"]),
              ("en double dans le meme onglet", totaux["doublon_jour"]),
              ("numero tronque par Excel", totaux["tronque"]),
              ("numero illisible", totaux["illisible"]),
              ("sans numero", totaux["vide"])]
    print(f"INUTILISABLE {livrees - nettes:>6}   sur {livrees} lignes livrees")
    for libelle, nombre in causes:
        if nombre:
            print(f"             {nombre:>6}   {libelle}")
    if totaux["etranger"]:
        print(f"\nDont hors Senegal {totaux['etranger']}, comptes utilisables.")
    if canaux:
        print("\nCanaux les plus frequents :")
        for canal, nombre in canaux.most_common(8):
            print(f"  {nombre:>6}  {canal}")


def verifier():
    cas = [("p:+221779103916", ("221779103916", "ok")), ("775618088", ("221775618088", "ok")),
           ("2.21773E+11", ("221773000000", "tronque")), ("", ("", "vide")),
           ("+33660583813", ("33660583813", "etranger")), ("22177", ("22177", "illisible")),
           ("Téléphone", ("", "illisible"))]
    for brut, attendu in cas:
        assert normaliser_telephone(brut) == attendu, (brut, normaliser_telephone(brut), attendu)
    print("ok")


def main(argv):
    if "--verifier" in argv:
        verifier()
        return 0
    chemins = [a for a in argv[1:] if not a.startswith("-")]
    sortie_csv = argv[argv.index("--csv") + 1] if "--csv" in argv else None
    if not chemins:
        print(__doc__)
        return 1
    feuilles = []
    for chemin in chemins:
        feuilles += lire_xlsx(chemin) if chemin.lower().endswith((".xlsx", ".xlsm")) else lire_csv(chemin)
    rapport, canaux, uniques = auditer(feuilles)
    if not rapport:
        print("Aucun onglet avec une colonne Telephone.")
        return 1
    afficher(rapport, canaux, uniques)
    if sortie_csv:
        with open(sortie_csv, "w", newline="", encoding="utf-8-sig") as f:
            ecrivain = csv.writer(f, delimiter=";")
            ecrivain.writerow(["onglet"] + [t for _, t in COLONNES_RAPPORT])
            for nom, jour in rapport.items():
                ecrivain.writerow([nom] + [jour[c] for c, _ in COLONNES_RAPPORT])
        print(f"\nRapport ecrit dans {sortie_csv}")
    return 0


if __name__ == "__main__":
    sys.exit(main(sys.argv))
